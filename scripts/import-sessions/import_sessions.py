#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convierte el Excel de sesiones (una fila por hora) en SQL para precargar reservaciones.

NO toca ninguna base de datos: solo lee el Excel y escribe dos archivos SQL que tú revisas y aplicas
(ver docs/DATA_MIGRATION.md):

    import.sql    carga todo en UNA transacción; es idempotente (correrlo dos veces no duplica nada)
    rollback.sql  deshace exactamente lo que cargó import.sql

Qué hace con el archivo:
  1. Valida (columnas, horas, domingos, traslapes, cierres...) y se detiene si hay errores.
  2. Junta las horas consecutivas de la misma persona/materia/día en UNA reservación.
  3. Detecta series semanales (>= --min-series repeticiones, mismo día de la semana y horario)
     y las guarda como reservaciones recurrentes (huecos permitidos, p. ej. un festivo).
  4. Crea las cuentas de las personas (rol académico, contraseña aleatoria inutilizable: cada
     persona la define con «¿Olvidaste tu contraseña?» o un administrador se la asigna).
  No envía correos (escribe directo a la base) y registra cada alta en el historial de cambios.

Dependencias (solo para esta herramienta):  pip install openpyxl bcrypt
"""
import argparse
import collections
import csv
import datetime as dt
import json
import os
import re
import secrets
import sys
import unicodedata
import uuid

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl:  pip install openpyxl bcrypt")

NS = uuid.UUID("5b0f3c3e-6a52-4f0e-9d0a-1b6a1c0d1e11")   # espacio de nombres fijo -> IDs deterministas (idempotencia)
TZ = "America/Mexico_City"
EXPECTED = ["fecha", "inicio", "fin", "sala", "materia", "nombre", "apaterno", "amaterno"]
SEEDED_ROOM_ID = "a10e1300-0000-4000-8000-000000000001"
MARKER_SOURCE = "import-sessions"


# ───────────────────────── utilidades ─────────────────────────
def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip())


def key_name(nombre, ap, am):
    f = lambda s: unicodedata.normalize("NFD", norm(s).upper()).encode("ascii", "ignore").decode()
    return (f(nombre), f(ap), f(am))


def q(v):
    """Literal SQL de texto."""
    return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"


def ts(date, time):
    return f"(TIMESTAMP '{date.isoformat()} {time.strftime('%H:%M:%S')}' AT TIME ZONE '{TZ}')"


def uid(*parts):
    return str(uuid.uuid5(NS, "|".join(map(str, parts))))


def hours(a, b):
    return (dt.datetime.combine(dt.date.today(), b) - dt.datetime.combine(dt.date.today(), a)).total_seconds() / 3600


# ───────────────────────── lectura y validación ─────────────────────────
def read_rows(path, errors):
    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [norm(c.value).lower() for c in ws[1]]
    if header[:8] != EXPECTED:
        errors.append(f"Encabezados inesperados: {header[:8]} (se esperaba {EXPECTED})")
        return []
    rows = []
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v in (None, "") for v in r):
            continue
        fecha, ini, fin, sala, materia, nombre, ap, am = (r + (None,) * 8)[:8]
        if isinstance(fecha, dt.datetime):
            fecha = fecha.date()
        if isinstance(ini, dt.datetime):
            ini = ini.time()
        if isinstance(fin, dt.datetime):
            fin = fin.time()
        bad = [name for name, v in zip(EXPECTED, (fecha, ini, fin, sala, materia, nombre, ap, am)) if v in (None, "")]
        if bad:
            errors.append(f"Fila {n}: faltan datos en {bad}")
            continue
        if not (isinstance(fecha, dt.date) and isinstance(ini, dt.time) and isinstance(fin, dt.time)):
            errors.append(f"Fila {n}: fecha u hora con formato inválido ({fecha!r}, {ini!r}, {fin!r})")
            continue
        if fin <= ini:
            errors.append(f"Fila {n}: la hora de fin ({fin}) no es posterior al inicio ({ini})")
            continue
        rows.append(dict(n=n, date=fecha, start=ini, end=fin, room=norm(sala), subject=norm(materia),
                         who=key_name(nombre, ap, am), raw=(norm(nombre), norm(ap), norm(am))))
    return rows


def merge_consecutive(rows):
    out = []
    for r in sorted(rows, key=lambda r: (r["date"], r["who"], r["subject"], r["start"])):
        p = out[-1] if out else None
        if p and (p["date"], p["who"], p["subject"], p["room"]) == (r["date"], r["who"], r["subject"], r["room"]) and p["end"] == r["start"]:
            p["end"] = r["end"]
            p["rows"].append(r["n"])
        else:
            out.append(dict(r, rows=[r["n"]]))
    return out


def detect_series(bookings, min_series):
    groups = collections.defaultdict(list)
    for b in bookings:
        groups[(b["who"], b["subject"], b["room"], b["date"].weekday(), b["start"], b["end"])].append(b)
    series = []
    for k, items in groups.items():
        items.sort(key=lambda b: b["date"])
        if len(items) >= min_series and all((items[i + 1]["date"] - items[i]["date"]).days % 7 == 0 for i in range(len(items) - 1)):
            series.append(items)
    for items in series:
        gid = uid("series", items[0]["who"], items[0]["subject"], items[0]["date"].weekday(), items[0]["start"], items[0]["end"])
        for b in items:
            b["group"] = gid
        items[0]["_series_info"] = dict(gid=gid, first=items[0]["date"], last=items[-1]["date"], count=len(items))
    return series


# ───────────────────────── programa principal ─────────────────────────
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="Excel de sesiones")
    ap.add_argument("--people", help="CSV con los correos de las personas (ver --init-people)")
    ap.add_argument("--init-people", metavar="CSV", help="Escribe una plantilla de personas (con correos vacíos) y termina")
    ap.add_argument("--room-name", help="Nombre real de la sala del archivo (renombra la sala sembrada)")
    ap.add_argument("--room-id", default=SEEDED_ROOM_ID, help="Id de la sala destino (por defecto, la sala sembrada)")
    ap.add_argument("--admin-email", help="Correo del usuario que figurará como creador de las reservaciones")
    ap.add_argument("--out-dir", default=".", help="Carpeta de salida de import.sql y rollback.sql")
    ap.add_argument("--min-series", type=int, default=3, help="Mínimo de repeticiones semanales para considerarlo serie (3)")
    ap.add_argument("--no-series", action="store_true", help="No crear series recurrentes")
    ap.add_argument("--no-merge", action="store_true", help="No juntar horas consecutivas (una reservación por fila)")
    ap.add_argument("--closure-dates", default="", help="Fechas de cierre separadas por coma (YYYY-MM-DD): si una reservación cae ahí, error")
    ap.add_argument("--holiday-dates", default="", help="Fechas festivas separadas por coma: solo advierte")
    ap.add_argument("--dry-run", action="store_true", help="Solo el reporte, sin escribir SQL")
    a = ap.parse_args()

    errors, warnings = [], []
    rows = read_rows(a.xlsx, errors)

    if a.init_people:
        seen = collections.OrderedDict()
        for r in rows:
            seen.setdefault(r["who"], r["raw"])
        with open(a.init_people, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["nombre", "apaterno", "amaterno", "correo", "nombre_mostrar"])
            for nom, ap_, am in seen.values():
                w.writerow([nom, ap_, am, "", ""])
        print(f"Plantilla escrita en {a.init_people}: llena 'correo' (obligatorio) y, si quieres acentos, 'nombre_mostrar'.")
        return

    # ── personas
    people = {}
    if a.people:
        with open(a.people, newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                k = key_name(r["nombre"], r["apaterno"], r["amaterno"])
                email = norm(r.get("correo")).lower()
                if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
                    errors.append(f"Persona {r['nombre']} {r['apaterno']}: correo inválido o vacío ({email!r})")
                    continue
                display = norm(r.get("nombre_mostrar")) or " ".join(norm(r[c]) for c in ("nombre", "apaterno", "amaterno")).title()
                people[k] = dict(email=email, name=display)
    emails = collections.Counter(p["email"] for p in people.values())
    errors += [f"Correo repetido en el CSV: {e}" for e, c in emails.items() if c > 1]
    for k in {r["who"] for r in rows}:
        if k not in people:
            errors.append(f"Falta el correo de {' '.join(k)} (agrégalo al CSV de personas)")

    # ── reglas de datos
    closures = {d.strip() for d in a.closure_dates.split(",") if d.strip()}
    holidays = {d.strip() for d in a.holiday_dates.split(",") if d.strip()}
    rooms = {r["room"] for r in rows}
    if len(rooms) > 1:
        errors.append(f"El archivo trae varias salas {sorted(rooms)}; esta herramienta carga una sola sala por corrida")
    for r in rows:
        if r["date"].weekday() == 6:
            errors.append(f"Fila {r['n']}: {r['date']} es domingo (no se reserva)")
        if not (dt.time(7, 0) <= r["start"] and r["end"] <= dt.time(21, 0)):
            errors.append(f"Fila {r['n']}: {r['start']}–{r['end']} queda fuera del horario 07:00–21:00")
        if r["date"].isoformat() in closures:
            errors.append(f"Fila {r['n']}: {r['date']} es un cierre institucional")
        if r["date"].isoformat() in holidays:
            w = f"{r['date']} es festivo (se puede reservar; confirma que es correcto)"
            if w not in warnings:
                warnings.append(w)
        if r["date"].weekday() == 5:
            w = f"{r['date']} es sábado"
            if w not in warnings:
                warnings.append(w)
    dup = [k for k, c in collections.Counter((r["date"], r["start"], r["end"], r["room"]) for r in rows).items() if c > 1]
    errors += [f"Franja repetida en el archivo: {d} {s}-{e}" for d, s, e, _ in dup]

    bookings = rows if a.no_merge else merge_consecutive(rows)
    if a.no_merge:
        for b in bookings:
            b["rows"] = [b["n"]]
    for i, x in enumerate(bookings):                        # traslapes dentro del archivo
        for y in bookings[i + 1:]:
            if x["room"] == y["room"] and x["date"] == y["date"] and x["start"] < y["end"] and y["start"] < x["end"]:
                errors.append(f"Traslape el {x['date']}: {x['start']}–{x['end']} ({' '.join(x['who'])}) con {y['start']}–{y['end']} ({' '.join(y['who'])})")
    series = [] if a.no_series else detect_series(bookings, a.min_series)

    # ── reporte
    print(f"Archivo: {len(rows)} filas  ->  {len(bookings)} reservaciones"
          f"{' (sin juntar horas)' if a.no_merge else ' (horas consecutivas juntadas)'}")
    by_len = collections.Counter(int(hours(b["start"], b["end"])) for b in bookings)
    print("  por duración:", ", ".join(f"{h} h: {c}" for h, c in sorted(by_len.items())))
    per = collections.Counter(b["who"] for b in bookings)
    for k, c in per.most_common():
        p = people.get(k)
        print(f"  {(p['name'] if p else ' '.join(k)):40} {c:3} reservaciones   {p['email'] if p else '(sin correo)'}")
    print(f"Series semanales: {len(series)}")
    for items in series:
        b0, bn = items[0], items[-1]
        print(f"  {' '.join(b0['who']):32} {b0['date'].strftime('%A'):9} {b0['start']:%H:%M}-{b0['end']:%H:%M}  x{len(items)}  {b0['date']} .. {bn['date']}")
    print(f"Reservaciones sueltas: {sum(1 for b in bookings if 'group' not in b)}")
    print(f"Rango de fechas: {min(b['date'] for b in bookings)} .. {max(b['date'] for b in bookings)}")
    for w in warnings:
        print("  AVISO:", w)
    if errors:
        print("\nERRORES (no se generó nada):")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    if a.dry_run:
        print("\n(dry-run: no se escribió SQL)")
        return
    if not (a.room_name and a.admin_email):
        sys.exit("Faltan --room-name y --admin-email para generar el SQL.")
    try:
        import bcrypt
    except ImportError:
        sys.exit("Falta bcrypt:  pip install bcrypt")

    # ── SQL
    room_id, admin = a.room_id, a.admin_email.lower()
    now = dt.date.today().isoformat()
    L = ["-- Generado por scripts/import-sessions/import_sessions.py (" + now + ")",
         f"-- {len(bookings)} reservaciones, {len(series)} series, {len(people)} personas. Idempotente y en una sola transacción.",
         "BEGIN;", "",
         "DO $$ BEGIN",
         f"  IF NOT EXISTS (SELECT 1 FROM rooms WHERE id = {q(room_id)}) THEN RAISE EXCEPTION 'La sala % no existe (¿corriste el despliegue con la migración 009?)', {q(room_id)}; END IF;",
         f"  IF NOT EXISTS (SELECT 1 FROM users WHERE email = {q(admin)}) THEN RAISE EXCEPTION 'No existe el usuario administrador %', {q(admin)}; END IF;",
         "END $$;", "",
         f"UPDATE rooms SET name = {q(a.room_name)} WHERE id = {q(room_id)} AND name <> {q(a.room_name)};", ""]
    user_ids = {}
    for k, p in people.items():
        if k not in {b["who"] for b in bookings}:
            continue
        user_ids[k] = uid("user", p["email"])
        pw = bcrypt.hashpw(secrets.token_urlsafe(32).encode(), bcrypt.gensalt(10)).decode()
        L.append(f"INSERT INTO users (id, name, email, password_hash, role, is_admin, active) VALUES "
                 f"({q(user_ids[k])}, {q(p['name'])}, {q(p['email'])}, {q(pw)}, 'academico', false, true) ON CONFLICT (email) DO NOTHING;")
    L.append("")
    for items in series:
        s = items[0]["_series_info"]
        L.append(f"INSERT INTO recurring_groups (id, pattern, end_date, max_occurrences) VALUES "
                 f"({q(s['gid'])}, 'weekly', DATE '{s['last'].isoformat()}', {s['count']}) ON CONFLICT (id) DO NOTHING;")
    L.append("")
    res_ids = []
    for b in sorted(bookings, key=lambda b: (b["date"], b["start"])):
        rid = uid("reservation", room_id, b["date"], b["start"], b["who"])
        res_ids.append(rid)
        email = people[b["who"]]["email"]
        L.append(
            "INSERT INTO reservations (id, responsible_id, responsible_name, area, start_time, end_time, status, "
            "is_recurring, recurring_group, created_by, last_modified_by, room_id) VALUES ("
            f"{q(rid)}, (SELECT id FROM users WHERE email = {q(email)}), {q(people[b['who']]['name'])}, {q(b['subject'])}, "
            f"{ts(b['date'], b['start'])}, {ts(b['date'], b['end'])}, 'active', "
            f"{'true' if 'group' in b else 'false'}, {q(b['group']) if 'group' in b else 'NULL'}, "
            f"(SELECT id FROM users WHERE email = {q(admin)}), (SELECT id FROM users WHERE email = {q(admin)}), {q(room_id)}) "
            "ON CONFLICT (id) DO NOTHING;")
    ids_sql = ", ".join(q(i) for i in res_ids)
    L += ["",
          "-- Historial de cambios: una entrada «Reservación creada» por reservación nueva",
          "INSERT INTO audit_log (user_id, action, entity, entity_id, details)",
          f"SELECT (SELECT id FROM users WHERE email = {q(admin)}), 'create_reservation', 'reservations', r.id,",
          f"       jsonb_build_object('responsible_name', r.responsible_name, 'room_name', rm.name, 'area', r.area,",
          f"                          'start_time', r.start_time, 'end_time', r.end_time, 'source', '{MARKER_SOURCE}')",
          "FROM reservations r JOIN rooms rm ON rm.id = r.room_id",
          f"WHERE r.id IN ({ids_sql})",
          "  AND NOT EXISTS (SELECT 1 FROM audit_log al WHERE al.entity_id = r.id AND al.action = 'create_reservation');", "",
          "-- Seguridad: nada cargado puede traslaparse con otra reservación activa de la misma sala",
          "DO $$ DECLARE n int; BEGIN",
          "  SELECT count(*) INTO n FROM reservations a JOIN reservations b",
          "    ON a.id < b.id AND a.room_id = b.room_id AND a.status = 'active' AND b.status = 'active'",
          "   AND a.start_time < b.end_time AND a.end_time > b.start_time",
          f"  WHERE a.id IN ({ids_sql}) OR b.id IN ({ids_sql});",
          "  IF n > 0 THEN RAISE EXCEPTION 'Hay % traslapes con reservaciones existentes: se cancela toda la carga', n; END IF;",
          "END $$;", "",
          f"SELECT count(*) AS reservaciones_cargadas FROM reservations WHERE id IN ({ids_sql});",
          "COMMIT;"]
    R = ["-- rollback de import.sql: borra solo lo que esa carga creó (por ID determinista)", "BEGIN;",
         f"DELETE FROM audit_log WHERE entity_id IN ({ids_sql});",
         f"DELETE FROM reservations WHERE id IN ({ids_sql});"]
    gids = ", ".join(q(items[0]['_series_info']['gid']) for items in series)
    if gids:
        R.append(f"DELETE FROM recurring_groups WHERE id IN ({gids}) AND NOT EXISTS (SELECT 1 FROM reservations WHERE recurring_group = recurring_groups.id);")
    uids = ", ".join(q(i) for i in user_ids.values())
    if uids:
        R.append(f"DELETE FROM users WHERE id IN ({uids}) AND NOT EXISTS (SELECT 1 FROM reservations WHERE responsible_id = users.id);")
    R += ["-- (el nombre de la sala NO se revierte: edítalo en Administración > Salas si hace falta)", "COMMIT;"]

    os.makedirs(a.out_dir, exist_ok=True)
    for name, lines in (("import.sql", L), ("rollback.sql", R)):
        with open(os.path.join(a.out_dir, name), "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    print(f"\nEscrito: {os.path.join(a.out_dir, 'import.sql')} y rollback.sql  (contienen datos personales: NO los subas a git)")


if __name__ == "__main__":
    main()
